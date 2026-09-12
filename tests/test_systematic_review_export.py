import uuid
from datetime import datetime, timezone

from openpyxl import load_workbook

from common.systematic_review_export import export_systematic_review_xlsx
from db.db import get_conn, log_review_audit


def _create_review(project_id):
    review_id = f"review-export-{uuid.uuid4().hex}"
    article_id = f"article-export-{uuid.uuid4().hex}"

    with get_conn() as conn, conn.cursor() as cur:
        cur.execute(
            """
            INSERT INTO systematic_reviews (
                id,
                project_id,
                title,
                review_type,
                research_question,
                population,
                intervention,
                comparator,
                outcomes,
                inclusion_criteria,
                exclusion_criteria,
                status,
                created_at,
                updated_at
            )
            VALUES (
                %s, %s, %s, %s, %s,
                %s, %s, %s, %s,
                %s, %s, %s, %s, %s
            )
            """,
            (
                review_id,
                project_id,
                "Systematic review export test",
                "systematic_review",
                "Does the intervention improve outcomes?",
                "Adults",
                "Intervention",
                "Comparator",
                "Clinical outcome",
                "Eligible studies",
                "Ineligible studies",
                "active",
                datetime.now(timezone.utc),
                datetime.now(timezone.utc),
            ),
        )

        cur.execute(
            """
            INSERT INTO review_articles (
                id,
                review_id,
                title,
                abstract,
                journal,
                year,
                screening_status,
                screening_stage,
                title_abstract_status,
                full_text_status,
                full_text_available,
                full_text_retrieval_status,
                final_decision,
                is_duplicate,
                created_at,
                updated_at
            )
            VALUES (
                %s, %s, %s, %s, %s, %s,
                'reviewed',
                'full_text',
                'include',
                'include',
                TRUE,
                'retrieved',
                'include',
                FALSE,
                %s,
                %s
            )
            """,
            (
                article_id,
                review_id,
                "Export test article",
                "Example abstract",
                "Example Journal",
                "2026",
                datetime.now(timezone.utc),
                datetime.now(timezone.utc),
            ),
        )

    log_review_audit(
        project_id=project_id,
        review_id=review_id,
        article_id=article_id,
        action="test_export_audit_event",
        actor_type="ai",
        actor_id="test-model",
        stage="data_extraction",
        after_state={
            "ai_value": 120,
            "ai_confidence": 0.95,
        },
        model="test-model",
        provider="test-provider",
        prompt_version="test-prompt-v1",
        details={
            "field_key": "sample_size",
            "quote_verified": True,
        },
    )

    return review_id


def _cleanup_review(review_id):
    with get_conn() as conn, conn.cursor() as cur:
        cur.execute(
            "DELETE FROM systematic_reviews WHERE id = %s",
            (review_id,),
        )


def test_systematic_review_excel_export(project_id):
    review_id = _create_review(project_id)

    try:
        path = export_systematic_review_xlsx(
            review_id=review_id,
            project_id=project_id,
        )

        assert path.exists()
        assert path.suffix == ".xlsx"
        assert path.stat().st_size > 0

        workbook = load_workbook(path)

        assert workbook.sheetnames == [
            "Resumen",
            "PRISMA",
            "Articulos",
            "Screening",
            "Conflictos",
            "Busquedas",
            "Estrategias",
            "Extracciones",
            "Campos_extraccion",
            "Auditoria",
        ]

        resumen = workbook["Resumen"]
        assert resumen["B3"].value == review_id
        assert resumen["B4"].value == "Systematic review export test"

        articulos = workbook["Articulos"]

        headers = [
            cell.value
            for cell in articulos[1]
        ]

        assert "title" in headers
        assert "full_text_retrieval_status" in headers
        assert "final_decision" in headers

        title_col = headers.index("title") + 1

        assert (
            articulos.cell(
                row=2,
                column=title_col,
            ).value
            == "Export test article"
        )

        auditoria = workbook["Auditoria"]

        audit_headers = [
            cell.value
            for cell in auditoria[1]
        ]

        assert "action" in audit_headers
        assert "actor_type" in audit_headers
        assert "model" in audit_headers
        assert "provider" in audit_headers
        assert "prompt_version" in audit_headers
        assert "after_state" in audit_headers
        assert "details" in audit_headers
        assert "created_at" in audit_headers

        action_col = audit_headers.index("action") + 1
        actor_type_col = audit_headers.index("actor_type") + 1
        model_col = audit_headers.index("model") + 1
        provider_col = audit_headers.index("provider") + 1
        prompt_col = audit_headers.index("prompt_version") + 1

        assert (
            auditoria.cell(
                row=2,
                column=action_col,
            ).value
            == "test_export_audit_event"
        )

        assert (
            auditoria.cell(
                row=2,
                column=actor_type_col,
            ).value
            == "ai"
        )

        assert (
            auditoria.cell(
                row=2,
                column=model_col,
            ).value
            == "test-model"
        )

        assert (
            auditoria.cell(
                row=2,
                column=provider_col,
            ).value
            == "test-provider"
        )

        assert (
            auditoria.cell(
                row=2,
                column=prompt_col,
            ).value
            == "test-prompt-v1"
        )

        # Regression test:
        # openpyxl cannot save timezone-aware datetimes.
        for sheet in workbook.worksheets:
            for row in sheet.iter_rows():
                for cell in row:
                    value = cell.value
                    if isinstance(value, datetime):
                        assert value.tzinfo is None

    finally:
        _cleanup_review(review_id)


def test_systematic_review_export_rejects_other_project(project_id):
    review_id = _create_review(project_id)

    try:
        try:
            export_systematic_review_xlsx(
                review_id=review_id,
                project_id="another-project",
            )
        except ValueError as exc:
            assert "proyecto activo" in str(exc)
        else:
            raise AssertionError(
                "Expected ValueError for wrong project"
            )

    finally:
        _cleanup_review(review_id)
