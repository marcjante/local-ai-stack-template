"""Real PostgreSQL/RQ coverage for multi-file ingestion and shared documents."""

import hashlib
import importlib.util
import io
from pathlib import Path
import uuid

from alembic.migration import MigrationContext
from alembic.operations import Operations
from docx import Document
from openpyxl import Workbook
from pypdf import PdfWriter
from pypdf.generic import DictionaryObject, NameObject, DecodedStreamObject
import psycopg2
import pytest
from rq.job import Job
import sqlalchemy as sa

from common import document_storage as storage, document_service as documents, thesis_files
from common.thesis_service import create_thesis
from db.db import get_conn, get_document, register_document, create_project, delete_document
from rag.retrieval import index_document
from workers.queue_conn import redis_conn


@pytest.fixture(autouse=True)
def isolated_originals(tmp_path, monkeypatch):
    monkeypatch.setattr(storage, "ROOT", tmp_path / "originals")


@pytest.fixture
def thesis_project(project_id):
    create_thesis(project_id, {})
    yield project_id
    with get_conn() as conn, conn.cursor() as cur:
        cur.execute("DELETE FROM thesis_file_ingestions WHERE project_id=%s", (project_id,))
        cur.execute("DELETE FROM systematic_reviews WHERE project_id=%s", (project_id,))
        cur.execute("DELETE FROM rag_chunks WHERE doc_id IN (SELECT doc_id FROM documents WHERE project_id=%s)", (project_id,))


def headers(subject="admin", role="admin"):
    from backend.auth import issue_token
    return {"Authorization": "Bearer " + issue_token(subject, role)}


def url(project_id, suffix=""):
    return f"/api/projects/{project_id}/thesis/files{suffix}"


def sample(ext):
    output = io.BytesIO()
    if ext == "csv":
        return b"name,value\nAna,12\nLuis,24\n"
    if ext == "docx":
        doc = Document()
        doc.add_heading("Metodología", level=1)
        doc.add_paragraph("Participaron doce personas en el estudio.")
        doc.save(output)
    elif ext == "xlsx":
        book = Workbook()
        sheet = book.active
        sheet.title = "Resultados"
        sheet.append(["nombre", "valor"])
        sheet.append(["Ana", 12])
        book.save(output)
        book.close()
    elif ext == "pdf":
        writer = PdfWriter()
        page = writer.add_blank_page(width=612, height=792)
        font = DictionaryObject({NameObject("/Type"): NameObject("/Font"),
                                 NameObject("/Subtype"): NameObject("/Type1"),
                                 NameObject("/BaseFont"): NameObject("/Helvetica")})
        page[NameObject("/Resources")] = DictionaryObject({NameObject("/Font"): DictionaryObject({NameObject("/F1"): font})})
        stream = DecodedStreamObject()
        stream.set_data(b"BT /F1 12 Tf 50 700 Td (Methods) Tj 0 -20 Td (Twelve participants.) Tj ET")
        page[NameObject("/Contents")] = stream
        writer.write(output)
    return output.getvalue()


def upload(client, pid, entries, **fields):
    return client.post(url(pid), data={"files[]": [(io.BytesIO(data), name) for name, data in entries], **fields},
                       headers=headers(), content_type="multipart/form-data")


def process(file):
    # Execute the jobs stored in real Redis, including the ingest -> extract dispatch.
    Job.fetch(file["job_id"], connection=redis_conn).perform()
    return Job.fetch(file["job_id"] + "-extract", connection=redis_conn).perform()


def test_multi_format_batch_and_locators(backend_client, thesis_project):
    entries = [("sample." + ext, sample(ext)) for ext in ("pdf", "docx", "xlsx", "csv")]
    response = upload(backend_client, thesis_project, entries, declared_role="own_data")
    assert response.status_code == 202
    data = response.json
    assert data["batch_id"] and len(data["accepted"]) == 4
    assert not data["rejected"] and not data["reused"]
    for file, (name, content) in zip(data["accepted"], entries):
        assert file["source_role"] == "own_data"
        assert file["content_hash"] == hashlib.sha256(content).hexdigest()
        doc = get_document(file["doc_id"])
        assert not Path(doc["storage_key"]).is_absolute()
        assert storage.path_for(doc["storage_key"]).read_bytes() == content
        assert name not in doc["storage_key"]
        process(file)
        detail = backend_client.get(url(thesis_project, "/" + file["id"]), headers=headers()).json
        assert detail["status"] == "indexed" and detail["fragment_count"] > 0
        assert detail["binary_identity"] == "sha256_verified"
        for fragment in detail["fragments"]:
            assert fragment["project_id"] == thesis_project and fragment["doc_id"] == file["doc_id"]
            assert fragment["fragment_hash"] == hashlib.sha256(fragment["text"].encode()).hexdigest()
            assert fragment["rag_chunk_id"]
            assert fragment["char_end"] > fragment["char_start"] >= 0
            locator = fragment["source_locator"]
            if name.endswith("pdf"):
                assert locator["page_number"] == 1 and locator["section"] == "Methods"
            elif name.endswith("docx"):
                assert locator["paragraph_index"] in (1, 2)
                assert "page_number" not in locator
            elif name.endswith("xlsx"):
                assert locator["sheet"] == "Resultados" and ":" in locator["range"]
            else:
                assert locator["columns"] == ["name", "value"]
                assert locator["line_end"] >= locator["line_start"] >= 2


@pytest.mark.parametrize("filename,content", [("bad.exe", b"MZbinary"), ("../../escape.csv", b"a,b\n1,2"),
                                              ("..\\escape.csv", b"a,b\n1,2"), ("fake.pdf", b"not a pdf"),
                                              ("fake.docx", b"not a zip"), ("binary.csv", b"a\x00b")])
def test_rejected_files_do_not_register(backend_client, thesis_project, filename, content):
    response = upload(backend_client, thesis_project, [(filename, content)])
    assert response.status_code == 400 and len(response.json["rejected"]) == 1
    assert thesis_files.list_files(thesis_project) == []
    assert not list(storage.ROOT.rglob("*.original"))


def test_wrong_mime(backend_client, thesis_project):
    response = backend_client.post(url(thesis_project), headers=headers(), data={
        "files[]": (io.BytesIO(sample("csv")), "a.csv", "image/png"),
    })
    assert response.status_code == 400 and response.json["rejected"]


def test_limit_and_partial_batch(backend_client, thesis_project, monkeypatch):
    monkeypatch.setenv("THESIS_MAX_FILE_MB", "1")
    response = upload(backend_client, thesis_project, [("large.csv", b"a" * (1024 * 1024 + 1)),
                                                       ("good.csv", sample("csv")), ("bad.exe", b"bad")])
    assert response.status_code == 202
    assert len(response.json["accepted"]) == 1 and len(response.json["rejected"]) == 2
    assert len(thesis_files.list_files(thesis_project)) == 1


def test_deduplication_reuses_chunks_and_embeddings(backend_client, thesis_project, monkeypatch):
    first = upload(backend_client, thesis_project, [("a.csv", sample("csv"))]).json["accepted"][0]
    process(first)
    before = thesis_files.file_detail(thesis_project, first["id"])
    def forbidden(*args, **kwargs):
        pytest.fail("duplicate extraction or embedding")
    monkeypatch.setattr(documents, "extract_structured_text", forbidden)
    monkeypatch.setattr(documents, "embed_text", forbidden)
    second = upload(backend_client, thesis_project, [("renamed.csv", sample("csv")), ("again.csv", sample("csv"))])
    assert len(second.json["reused"]) == 2 and not second.json["accepted"]
    assert {f["id"] for f in second.json["reused"]} == {first["id"]}
    from workers.plugins.thesis_extract import handle
    handle(first["job_id"], {"project_id": thesis_project, "file_id": first["id"]})
    after = thesis_files.file_detail(thesis_project, first["id"])
    assert after["fragments"] == before["fragments"]


def historical(pid):
    doc_id = str(uuid.uuid4())
    text = "Historical scientific evidence already extracted and indexed."
    count = index_document(doc_id, text)
    register_document(doc_id, "default", "history.pdf", "application/pdf", "v1", "hashing_trick_256", text, count, project_id=pid)
    return doc_id


def test_legacy_link_idempotent_without_hash_or_reextraction(backend_client, thesis_project, monkeypatch):
    doc_id = historical(thesis_project)
    def forbidden(*args, **kwargs):
        pytest.fail("legacy document must reuse existing chunks")
    monkeypatch.setattr(documents, "extract_structured_text", forbidden)
    monkeypatch.setattr(documents, "embed_text", forbidden)
    payload = {"doc_id": doc_id, "source_role": "scientific_evidence"}
    first = backend_client.post(url(thesis_project, "/link"), json=payload, headers=headers())
    second = backend_client.post(url(thesis_project, "/link"), json=payload, headers=headers())
    assert first.status_code == 201 and second.status_code == 200
    assert first.json["id"] == second.json["id"]
    assert first.json["fragments"] == second.json["fragments"]
    assert first.json["binary_identity"] == "legacy_unverified"
    assert first.json["content_hash"] is None and not first.json["original_available"]
    assert get_document(doc_id)["content_hash"] is None
    assert first.json["fragment_count"] == 1


def test_projects_are_independent_and_cross_access_is_rejected(backend_client, thesis_project):
    other = "test-other-" + uuid.uuid4().hex
    create_project(other, "Other")
    create_thesis(other, {})
    try:
        a = upload(backend_client, thesis_project, [("a.csv", sample("csv"))]).json["accepted"][0]
        b = upload(backend_client, other, [("a.csv", sample("csv"))]).json["accepted"][0]
        assert a["doc_id"] != b["doc_id"] and a["content_hash"] == b["content_hash"]
        assert get_document(a["doc_id"])["storage_key"] != get_document(b["doc_id"])["storage_key"]
        assert backend_client.get(url(thesis_project, "/" + b["id"]), headers=headers()).status_code == 404
        assert backend_client.post(url(thesis_project, "/link"), json={"doc_id": b["doc_id"]}, headers=headers()).status_code == 404
        listing = backend_client.get(url(thesis_project), headers=headers()).json["files"]
        assert [f["id"] for f in listing] == [a["id"]]
    finally:
        with get_conn() as conn, conn.cursor() as cur:
            cur.execute("DELETE FROM thesis_file_ingestions WHERE project_id=%s", (other,))
            cur.execute("DELETE FROM tasks WHERE project_id=%s", (other,))
            cur.execute("DELETE FROM projects WHERE id=%s", (other,))


@pytest.mark.parametrize("method,suffix", [("post", ""), ("get", ""), ("get", "/missing"), ("post", "/link")])
def test_authentication_and_nonexistent_project(backend_client, thesis_project, method, suffix):
    action = getattr(backend_client, method)
    assert action(url(thesis_project, suffix), json={}).status_code == 401
    assert action(url("missing-" + uuid.uuid4().hex, suffix), json={}, headers=headers()).status_code == 404
    assert action(url(thesis_project, suffix), json={}, headers=headers("outsider", "lector")).status_code == 403


def test_extraction_retry_is_atomic(backend_client, thesis_project, monkeypatch):
    first = upload(backend_client, thesis_project, [("a.csv", sample("csv"))]).json["accepted"][0]
    real = documents.extract_structured_text
    def fail(*args):
        raise ValueError("test extraction error")
    monkeypatch.setattr(documents, "extract_structured_text", fail)
    with pytest.raises(ValueError, match="test extraction error"):
        process(first)
    failed = thesis_files.file_detail(thesis_project, first["id"])
    assert failed["status"] == "failed" and failed["fragment_count"] == 0
    assert "test extraction error" in failed["error"]
    assert get_document(first["doc_id"])["n_chunks"] == 0
    monkeypatch.setattr(documents, "extract_structured_text", real)
    retry = upload(backend_client, thesis_project, [("a.csv", sample("csv"))]).json["reused"][0]
    assert retry["id"] == first["id"] and retry["job_id"] != first["job_id"]
    process(retry)
    assert thesis_files.file_detail(thesis_project, first["id"])["status"] == "indexed"


def test_atomic_storage_failure(backend_client, thesis_project, monkeypatch):
    def fail(*args):
        raise OSError("simulated disk failure before publish")
    monkeypatch.setattr(storage.os, "replace", fail)
    response = upload(backend_client, thesis_project, [("a.csv", sample("csv"))])
    assert response.status_code == 400
    assert not thesis_files.list_files(thesis_project)
    assert not [p for p in storage.ROOT.rglob("*") if p.is_file()]
    with get_conn() as conn, conn.cursor() as cur:
        cur.execute("SELECT count(*) FROM documents WHERE project_id=%s", (thesis_project,))
        assert cur.fetchone()[0] == 0


@pytest.mark.parametrize("linked", [False, True])
def test_review_replacement_preserves_shared_document(dashboard_client, backend_client, thesis_project, linked):
    from tests.test_systematic_review_full_text_upload import _create_review_with_article
    review_id, article_id = _create_review_with_article(thesis_project)
    def review_upload(data):
        return dashboard_client.post("/api/systematic-review/full-text-upload", data={
            "review_id": review_id, "article_id": article_id,
            "file": (io.BytesIO(data), "source.csv"),
        })
    first = review_upload(sample("csv"))
    assert first.status_code == 201
    old = first.json["document_id"]
    key = get_document(old)["storage_key"]
    if linked:
        result = backend_client.post(url(thesis_project, "/link"), json={"doc_id": old}, headers=headers())
        assert result.status_code == 201
        before = result.json["fragments"]
        assert delete_document(old) is False
    second = review_upload(b"name,value\nDifferent,42\n")
    assert second.status_code == 201 and second.json["document_id"] != old
    if linked:
        assert get_document(old) is not None
        assert storage.path_for(key).read_bytes() == sample("csv")
        assert backend_client.get(url(thesis_project, "/" + result.json["id"]), headers=headers()).json["fragments"] == before
        from db.db import get_review_audit_trail
        assert any(e["action"] == "shared_document_retained" for e in get_review_audit_trail(review_id, project_id=thesis_project, article_id=article_id))
        with pytest.raises(psycopg2.IntegrityError):
            with get_conn() as conn, conn.cursor() as cur:
                cur.execute("DELETE FROM documents WHERE doc_id=%s", (old,))
    else:
        assert get_document(old) is None


def test_review_document_reused_by_hash(dashboard_client, backend_client, thesis_project, monkeypatch):
    from tests.test_systematic_review_full_text_upload import _create_review_with_article
    review_id, article_id = _create_review_with_article(thesis_project)
    content = sample("pdf")
    original = dashboard_client.post("/api/systematic-review/full-text-upload", data={
        "review_id": review_id, "article_id": article_id, "file": (io.BytesIO(content), "source.pdf"),
    })
    assert original.status_code == 201
    def forbidden(*args, **kwargs):
        pytest.fail("shared evidence must not be indexed again")
    monkeypatch.setattr(documents, "extract_structured_text", forbidden)
    monkeypatch.setattr(documents, "embed_text", forbidden)
    reused = upload(backend_client, thesis_project, [("other-name.pdf", content)])
    assert reused.status_code == 202
    assert reused.json["reused"][0]["doc_id"] == original.json["document_id"]
    assert reused.json["reused"][0]["status"] == "indexed"


def test_migration_round_trip_and_nullable_legacy_hash():
    from db.db import DB_DSN
    path = Path("db/migrations/versions/b1b18d204a61_add_thesis_ingestion.py")
    spec = importlib.util.spec_from_file_location("thesis_ingestion_migration", path)
    migration = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(migration)
    assert migration.down_revision == "f1a7c93d820b"
    engine = sa.create_engine("postgresql+psycopg2://", creator=lambda: psycopg2.connect(DB_DSN))
    try:
        with engine.connect() as conn:
            tx = conn.begin()
            try:
                schema = "test_ingestion_" + uuid.uuid4().hex
                conn.exec_driver_sql(f'CREATE SCHEMA "{schema}"')
                conn.exec_driver_sql(f'SET LOCAL search_path TO "{schema}"')
                conn.exec_driver_sql("CREATE TABLE thesis_profiles (project_id TEXT PRIMARY KEY)")
                conn.exec_driver_sql("CREATE TABLE documents (doc_id TEXT PRIMARY KEY, project_id TEXT NOT NULL)")
                conn.exec_driver_sql("CREATE TABLE rag_chunks (chunk_id TEXT PRIMARY KEY, doc_id TEXT NOT NULL)")
                conn.exec_driver_sql("INSERT INTO documents VALUES ('legacy','p')")
                with Operations.context(MigrationContext.configure(conn)):
                    migration.upgrade()
                    assert conn.exec_driver_sql("SELECT content_hash FROM documents WHERE doc_id='legacy'").scalar() is None
                    columns = {c["name"] for c in sa.inspect(conn).get_columns("thesis_fragments", schema=schema)}
                    assert "text" not in columns and "rag_chunk_id" in columns
                    assert any(fk["referred_table"] == "documents" for fk in sa.inspect(conn).get_foreign_keys("thesis_file_ingestions", schema=schema))
                    migration.downgrade()
                    assert conn.exec_driver_sql("SELECT doc_id FROM documents").scalar() == "legacy"
                    assert "thesis_fragments" not in sa.inspect(conn).get_table_names(schema=schema)
                    migration.upgrade()
            finally:
                tx.rollback()
    finally:
        engine.dispose()
