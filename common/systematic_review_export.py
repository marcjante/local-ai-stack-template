"""Excel export for systematic reviews."""

import json
import tempfile
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from db.db import get_conn
from dashboard.prisma import calculate_prisma
from psycopg2.extras import RealDictCursor


HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
HEADER_FONT = Font(color="FFFFFF", bold=True)
TITLE_FONT = Font(size=16, bold=True)


def _cell_value(value):
    """Convert database values into Excel-safe values."""
    if value is None:
        return ""

    if isinstance(value, datetime):
        if value.tzinfo is not None:
            value = value.replace(tzinfo=None)
        return value

    if isinstance(value, (dict, list, tuple)):
        return json.dumps(
            value,
            ensure_ascii=False,
            default=str,
        )

    return value


def _write_table(ws, headers, rows):
    ws.append(headers)

    for cell in ws[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(
            horizontal="center",
            vertical="center",
            wrap_text=True,
        )

    for row in rows:
        ws.append([
            _cell_value(row.get(header))
            for header in headers
        ])

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions

    for column_index, column_cells in enumerate(
        ws.columns,
        start=1,
    ):
        max_length = 0

        for cell in column_cells:
            value = "" if cell.value is None else str(cell.value)
            max_length = max(
                max_length,
                min(len(value), 60),
            )

            if cell.row > 1:
                cell.alignment = Alignment(
                    vertical="top",
                    wrap_text=True,
                )

        ws.column_dimensions[
            get_column_letter(column_index)
        ].width = max(12, min(max_length + 2, 60))


def _fetch_rows(cur, query, params):
    cur.execute(query, params)
    return [dict(row) for row in cur.fetchall()]


def export_systematic_review_xlsx(
    review_id: str,
    project_id: str,
) -> Path:
    """Create a formatted XLSX export for one review."""

    with get_conn() as conn, conn.cursor(
        cursor_factory=RealDictCursor
    ) as cur:

        cur.execute(
            """
            SELECT
                id,
                project_id,
                title,
                review_type,
                research_question,
                population,
                intervention,
                comparator,
                outcomes,
                inclusion_criteria,
                exclusion_criteria,
                status,
                created_at,
                updated_at
            FROM systematic_reviews
            WHERE id = %s
              AND project_id = %s
            """,
            (review_id, project_id),
        )

        review_row = cur.fetchone()

        if not review_row:
            raise ValueError(
                "Revisión no encontrada en el proyecto activo"
            )

        review = dict(review_row)

        articles = _fetch_rows(
            cur,
            """
            SELECT
                ra.id,
                ra.search_id,
                ra.pmid,
                ra.doi,
                ra.title,
                ra.abstract,
                ra.authors,
                ra.journal,
                ra.year,
                ra.screening_status,
                ra.screening_stage,
                ra.ai_decision,
                ra.ai_reason,
                ra.ai_confidence,
                ra.human_decision,
                ra.exclusion_reason,
                ra.exclusion_reason_code,
                ra.title_abstract_status,
                ra.full_text_status,
                ra.full_text_available,
                ra.full_text_retrieval_status,
                ra.full_text_document_id,
                d.filename AS full_text_filename,
                d.n_chunks AS full_text_n_chunks,
                ra.final_decision,
                ra.is_duplicate,
                ra.duplicate_of_article_id,
                ra.duplicate_reason,
                ra.created_at,
                ra.updated_at
            FROM review_articles ra
            LEFT JOIN documents d
                ON d.doc_id = ra.full_text_document_id
            WHERE ra.review_id = %s
            ORDER BY ra.created_at ASC
            """,
            (review_id,),
        )

        decisions = _fetch_rows(
            cur,
            """
            SELECT
                d.id,
                d.review_id,
                d.article_id,
                a.title AS article_title,
                d.stage,
                d.reviewer_id,
                d.decision,
                d.exclusion_reason_code,
                d.reason,
                d.notes,
                d.created_at,
                d.updated_at
            FROM review_screening_decisions d
            JOIN review_articles a
                ON a.id = d.article_id
            WHERE d.review_id = %s
            ORDER BY d.created_at ASC
            """,
            (review_id,),
        )

        conflicts = _fetch_rows(
            cur,
            """
            SELECT
                c.id,
                c.review_id,
                c.article_id,
                a.title AS article_title,
                c.stage,
                c.status,
                c.resolution,
                c.resolved_by,
                c.resolution_notes,
                c.created_at,
                c.resolved_at
            FROM review_screening_conflicts c
            JOIN review_articles a
                ON a.id = c.article_id
            WHERE c.review_id = %s
            ORDER BY c.created_at ASC
            """,
            (review_id,),
        )

        searches = _fetch_rows(
            cur,
            """
            SELECT
                id,
                review_id,
                database_name,
                query,
                total_found,
                imported_count,
                searched_at,
                strategy_id
            FROM review_searches
            WHERE review_id = %s
            ORDER BY searched_at ASC
            """,
            (review_id,),
        )

        strategies = _fetch_rows(
            cur,
            """
            SELECT
                id,
                review_id,
                database_name,
                version,
                query,
                is_valid,
                accepted_by_database,
                total_found,
                query_translation,
                warnings,
                errors,
                removed_terms,
                concepts,
                model,
                provider,
                human_confirmed,
                created_at,
                confirmed_at
            FROM review_search_strategies
            WHERE review_id = %s
            ORDER BY database_name ASC, version ASC
            """,
            (review_id,),
        )

        extraction_fields = _fetch_rows(
            cur,
            """
            SELECT
                id,
                review_id,
                field_key,
                label,
                description,
                value_type,
                required,
                display_order,
                created_at,
                updated_at
            FROM review_extraction_fields
            WHERE review_id = %s
            ORDER BY display_order ASC, label ASC
            """,
            (review_id,),
        )

        extractions = _fetch_rows(
            cur,
            """
            SELECT
                e.id,
                e.review_id,
                e.article_id,
                a.title AS article_title,
                a.pmid,
                a.doi,
                e.field_id,
                f.field_key,
                f.label AS field_label,
                e.ai_value,
                e.ai_reason,
                e.ai_confidence,
                e.source_type,
                e.source_location,
                e.source_quote,
                e.human_value,
                e.validation_status,
                e.reviewer_id,
                e.reviewer_notes,
                e.model,
                e.provider,
                e.skills_used,
                e.created_at,
                e.updated_at
            FROM review_extractions e
            JOIN review_articles a
                ON a.id = e.article_id
            JOIN review_extraction_fields f
                ON f.id = e.field_id
            WHERE e.review_id = %s
            ORDER BY
                a.title ASC,
                f.display_order ASC,
                f.label ASC
            """,
            (review_id,),
        )


        audit_rows = _fetch_rows(
            cur,
            """
            SELECT
                ral.id,
                ral.project_id,
                ral.review_id,
                ral.article_id,
                a.title AS article_title,
                ral.extraction_id,
                f.field_key,
                f.label AS field_label,
                ral.action,
                ral.actor_type,
                ral.actor_id,
                ral.stage,
                ral.before_state,
                ral.after_state,
                ral.model,
                ral.provider,
                ral.prompt_version,
                ral.details,
                ral.created_at
            FROM review_audit_log ral
            LEFT JOIN review_articles a
                ON a.id = ral.article_id
            LEFT JOIN review_extractions e
                ON e.id = ral.extraction_id
            LEFT JOIN review_extraction_fields f
                ON f.id = e.field_id
            WHERE ral.review_id = %s
              AND ral.project_id = %s
            ORDER BY ral.created_at ASC, ral.id ASC
            """,
            (
                review_id,
                project_id,
            ),
        )

    prisma = calculate_prisma(
        searches,
        articles,
    )

    wb = Workbook()
    default_ws = wb.active
    wb.remove(default_ws)

    # ------------------------------------------------------
    # Resumen
    # ------------------------------------------------------

    ws = wb.create_sheet("Resumen")
    ws["A1"] = "Revisión sistemática"
    ws["A1"].font = TITLE_FONT

    summary_rows = [
        ("ID", review["id"]),
        ("Título", review["title"]),
        ("Tipo", review["review_type"]),
        ("Pregunta de investigación", review["research_question"]),
        ("Población", review["population"]),
        ("Intervención", review["intervention"]),
        ("Comparador", review["comparator"]),
        ("Outcomes", review["outcomes"]),
        ("Criterios de inclusión", review["inclusion_criteria"]),
        ("Criterios de exclusión", review["exclusion_criteria"]),
        ("Estado", review["status"]),
        ("Creada", review["created_at"]),
        ("Actualizada", review["updated_at"]),
        ("Proyecto", review["project_id"]),
    ]

    for index, (label, value) in enumerate(
        summary_rows,
        start=3,
    ):
        ws.cell(index, 1, label).font = Font(bold=True)
        ws.cell(index, 2, _cell_value(value))
        ws.cell(index, 2).alignment = Alignment(
            vertical="top",
            wrap_text=True,
        )

    ws.column_dimensions["A"].width = 28
    ws.column_dimensions["B"].width = 90

    # ------------------------------------------------------
    # PRISMA
    # ------------------------------------------------------

    ws = wb.create_sheet("PRISMA")

    prisma_rows = [
        ("Registros identificados", prisma["records_identified"]),
        ("Registros importados", prisma["records_imported"]),
        ("Duplicados eliminados", prisma["duplicates_removed"]),
        ("Registros cribados", prisma["records_screened"]),
        ("Registros excluidos", prisma["records_excluded"]),
        ("Informes buscados", prisma["reports_sought"]),
        (
            "Informes no recuperados",
            prisma["reports_not_retrieved"],
        ),
        (
            "Textos completos evaluados",
            prisma["reports_assessed"],
        ),
        (
            "Textos completos excluidos",
            prisma["reports_excluded"],
        ),
        ("Estudios incluidos", prisma["studies_included"]),
    ]

    ws.append(["Indicador", "N"])

    for row in prisma_rows:
        ws.append(row)

    start_row = ws.max_row + 3
    ws.cell(start_row, 1, "Motivos de exclusión")
    ws.cell(start_row, 1).font = Font(bold=True)

    ws.cell(start_row + 1, 1, "Código")
    ws.cell(start_row + 1, 2, "N")

    current_row = start_row + 2

    for reason, count in prisma["exclusion_reasons"].items():
        ws.cell(current_row, 1, reason)
        ws.cell(current_row, 2, count)
        current_row += 1

    current_row += 1
    ws.cell(current_row, 1, "Bases de datos")
    ws.cell(current_row, 1).font = Font(bold=True)

    current_row += 1
    ws.cell(current_row, 1, "Base")
    ws.cell(current_row, 2, "Identificados")
    ws.cell(current_row, 3, "Importados")

    for database, values in prisma["databases"].items():
        current_row += 1
        ws.cell(current_row, 1, database)
        ws.cell(
            current_row,
            2,
            values["records_identified"],
        )
        ws.cell(
            current_row,
            3,
            values["records_imported"],
        )

    for cell in ws[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT

    ws.column_dimensions["A"].width = 40
    ws.column_dimensions["B"].width = 18
    ws.column_dimensions["C"].width = 18

    # ------------------------------------------------------
    # Tablas
    # ------------------------------------------------------

    table_specs = [
        ("Articulos", articles),
        ("Screening", decisions),
        ("Conflictos", conflicts),
        ("Busquedas", searches),
        ("Estrategias", strategies),
        ("Extracciones", extractions),
        ("Campos_extraccion", extraction_fields),
        ("Auditoria", audit_rows),
    ]

    for sheet_name, rows in table_specs:
        ws = wb.create_sheet(sheet_name)

        if rows:
            headers = list(rows[0].keys())
            _write_table(
                ws,
                headers,
                rows,
            )
        else:
            ws["A1"] = "Sin datos"
            ws["A1"].font = Font(bold=True)

    safe_title = "".join(
        char if char.isalnum() or char in "-_ " else "_"
        for char in review["title"]
    ).strip()

    if not safe_title:
        safe_title = "revision_sistematica"

    safe_title = safe_title[:70]

    tmp = tempfile.NamedTemporaryFile(
        prefix="systematic_review_",
        suffix=".xlsx",
        delete=False,
    )
    tmp.close()

    output_path = Path(tmp.name)

    wb.save(output_path)

    return output_path
